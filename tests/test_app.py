"""Tests for the Mock Data Factory app."""
from __future__ import annotations

import io
import json
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import (  # noqa: E402
    MAX_CELLS,
    MAX_CONTENT_LENGTH,
    MAX_FIELDS,
    MAX_ROWS,
    MAX_TEMPLATE_LENGTH,
    SchemaError,
    _infer_sql_type,
    _name_to_type,
    _sanitize_xml_name,
    _validate_schema,
    app,
    compile_template,
    format_csv,
    format_excel,
    format_json,
    format_sql,
    format_xml,
    generate_data,
    generate_value,
    infer_from_json,
    infer_from_sql,
    infer_from_typescript,
    render_template_value,
)


@pytest.fixture
def client():
    app.config["TESTING"] = True
    with app.test_client() as client:
        yield client


@pytest.fixture
def schema():
    return {
        "fields": [
            {"name": "id", "type": "Row Number"},
            {"name": "first_name", "type": "First Name"},
            {"name": "email", "type": "Email Address"},
        ],
        "num_rows": 5,
        "format": "JSON",
    }


# ---- generate_value -------------------------------------------------------

def test_generate_value_known_types_produce_values():
    for t in ["First Name", "Last Name", "Email Address", "City", "Country",
              "Number", "Decimal", "Date", "Gender", "IP Address v4", "Phone Number"]:
        assert generate_value(t) is not None


def test_generate_value_unknown_type_returns_empty_string():
    assert generate_value("Not A Real Type") == ""


def test_custom_list_uses_provided_values():
    config = {"values": ["a", "b", "c"]}
    for _ in range(20):
        assert generate_value("Custom List", config) in {"a", "b", "c"}


def test_custom_list_without_values_returns_empty():
    assert generate_value("Custom List", {}) == ""
    assert generate_value("Custom List", None) == ""


# ---- generate_data --------------------------------------------------------

def test_generate_data_respects_num_rows(schema):
    data = generate_data(schema)
    assert len(data) == 5


def test_generate_data_caps_at_max_rows():
    schema = {
        "fields": [{"name": "id", "type": "Row Number"}],
        "num_rows": MAX_ROWS + 100,
        "format": "JSON",
    }
    data = generate_data(schema)
    assert len(data) == MAX_ROWS


def test_generate_data_row_number_is_sequential(schema):
    data = generate_data(schema)
    assert [row["id"] for row in data] == [1, 2, 3, 4, 5]


def test_generate_data_preserves_field_order():
    schema = {
        "fields": [
            {"name": "z_field", "type": "First Name"},
            {"name": "a_field", "type": "Last Name"},
            {"name": "m_field", "type": "City"},
        ],
        "num_rows": 3,
        "format": "JSON",
    }
    data = generate_data(schema)
    for row in data:
        assert list(row.keys()) == ["z_field", "a_field", "m_field"]


def test_generate_data_rejects_zero_rows():
    schema = {"fields": [{"name": "id", "type": "Row Number"}], "num_rows": 0, "format": "JSON"}
    with pytest.raises(SchemaError):
        generate_data(schema)


def test_blank_percentage_modifier_produces_nulls():
    schema = {
        "fields": [{"name": "name", "type": "First Name", "blank_percentage": 100}],
        "num_rows": 10,
        "format": "JSON",
    }
    data = generate_data(schema)
    assert all(row["name"] is None for row in data)


def test_blank_null_type_is_empty_string_when_not_nulled():
    # Regression: Blank/Null used to fall back to a random fake word when a
    # row wasn't chosen for blanking, so the column was never truly blank.
    config = {"blank_percentage": 0}
    for _ in range(20):
        assert generate_value("Blank/Null", config) == ""


def test_blank_null_type_respects_blank_percentage():
    schema = {
        "fields": [{"name": "x", "type": "Blank/Null", "blank_percentage": 100}],
        "num_rows": 10,
        "format": "JSON",
    }
    data = generate_data(schema)
    assert all(row["x"] is None for row in data)


# ---- formatters -----------------------------------------------------------

def test_format_csv_writes_header_and_rows():
    data = [{"a": 1, "b": "x"}, {"a": 2, "b": "y"}]
    out = format_csv(data)
    lines = out.strip().splitlines()
    assert lines[0] == "a,b"
    assert lines[1] == "1,x"
    assert lines[2] == "2,y"


def test_format_csv_empty():
    assert format_csv([]) == ""


def test_format_csv_neutralizes_formula_injection():
    # Values that spreadsheet apps would interpret as formulas must be
    # prefixed with a single quote so they're imported as literal text.
    data = [{"a": "=1+1", "b": "+cmd", "c": "-x", "d": "@x", "e": "safe", "f": 5}]
    out = format_csv(data)
    lines = out.strip().splitlines()
    assert lines[1] == "'=1+1,'+cmd,'-x,'@x,safe,5"


def test_format_excel_neutralizes_formula_injection():
    data = [{"a": "=1+1", "b": "safe", "c": "+1-202-555-0100", "d": "#N/A", "e": 7}]
    xlsx_bytes = format_excel(data)
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    cell_a2 = ws.cell(row=2, column=1)
    # Stored as text, not as a formula ("f"), and without a visible quote prefix.
    assert cell_a2.value == "=1+1"
    assert cell_a2.data_type == "s"
    assert ws.cell(row=2, column=2).value == "safe"
    assert ws.cell(row=2, column=3).value == "+1-202-555-0100"
    assert ws.cell(row=2, column=4).data_type == "s"
    assert ws.cell(row=2, column=5).value == 7


def test_format_json_is_parseable():
    data = [{"x": 1}, {"x": 2}]
    parsed = json.loads(format_json(data))
    assert parsed == data


def test_format_xml_has_record_per_row():
    data = [{"name": "Alice"}, {"name": "Bob"}]
    out = format_xml(data)
    assert "<records>" in out
    assert out.count("<record>") == 2
    assert "Alice" in out and "Bob" in out


def test_format_xml_handles_none_as_empty():
    data = [{"x": None}]
    out = format_xml(data)
    assert "<x/>" in out or "<x></x>" in out


def test_format_sql_escapes_single_quotes():
    data = [{"name": "O'Hara"}]
    out = format_sql(data)
    assert "'O''Hara'" in out


def test_format_sql_uses_correct_column_types():
    data = [
        {"id": 1, "price": 9.99, "name": "x", "created": "2024-01-01"},
    ]
    out = format_sql(data)
    assert '"id" INTEGER' in out
    assert '"price" REAL' in out
    assert '"name" TEXT' in out
    assert '"created" DATE' in out


def test_infer_sql_type_skips_leading_nulls():
    # Bug-regression: original code broke on first iteration even if value was None.
    data = [{"a": None}, {"a": None}, {"a": 42}]
    assert _infer_sql_type("a", data) == "INTEGER"


def test_format_sql_handles_bool_as_integer():
    data = [{"flag": True}, {"flag": False}]
    out = format_sql(data)
    assert '"flag" INTEGER' in out
    assert "VALUES (1)" in out
    assert "VALUES (0)" in out


def test_format_sql_empty():
    assert format_sql([]) == ""


def test_format_sql_quotes_malicious_column_name():
    # Regression: a field name containing SQL syntax must not break out of
    # the identifier position in CREATE TABLE / INSERT.
    data = [{'evil"); DROP TABLE users; --': "x"}]
    out = format_sql(data)
    assert '"evil""); DROP TABLE users; --" TEXT' in out
    assert 'INSERT INTO "mock_data" ("evil""); DROP TABLE users; --")' in out


def test_format_sql_quotes_table_name():
    out = format_sql([{"a": 1}], table_name='t"); DROP TABLE users; --')
    assert 'CREATE TABLE IF NOT EXISTS "t""); DROP TABLE users; --" (' in out
    assert 'INSERT INTO "t""); DROP TABLE users; --"' in out


# ---- _validate_schema -----------------------------------------------------

def test_validate_schema_rejects_missing_keys():
    with pytest.raises(SchemaError):
        _validate_schema({"fields": []})


def test_validate_schema_rejects_empty_fields():
    with pytest.raises(SchemaError):
        _validate_schema({"fields": [], "num_rows": 1, "format": "CSV"})


def test_validate_schema_rejects_unknown_format():
    with pytest.raises(SchemaError):
        _validate_schema({"fields": [{"name": "a", "type": "First Name"}], "num_rows": 1, "format": "yaml"})


def test_validate_schema_rejects_negative_rows():
    with pytest.raises(SchemaError):
        _validate_schema({"fields": [{"name": "a", "type": "First Name"}], "num_rows": -1, "format": "CSV"})


def test_validate_schema_rejects_excessive_rows():
    with pytest.raises(SchemaError):
        _validate_schema({
            "fields": [{"name": "a", "type": "First Name"}],
            "num_rows": MAX_ROWS + 1,
            "format": "CSV",
        })


# ---- HTTP endpoints -------------------------------------------------------

def test_index_renders(client):
    response = client.get("/")
    assert response.status_code == 200
    assert b"Mock Data Factory" in response.data


def test_preview_returns_data_for_csv(client, schema):
    schema["format"] = "CSV"
    response = client.post("/preview", json=schema)
    assert response.status_code == 200
    payload = response.get_json()
    assert "data" in payload
    assert "field_order" in payload
    assert payload["field_order"] == ["id", "first_name", "email"]


def test_preview_returns_text_for_json(client, schema):
    response = client.post("/preview", json=schema)
    assert response.status_code == 200
    json.loads(response.get_data(as_text=True))  # must parse


def test_preview_rejects_invalid_schema(client):
    response = client.post("/preview", json={"bad": "schema"})
    assert response.status_code == 400


def test_preview_rejects_oversized_request(client):
    response = client.post("/preview", json={
        "fields": [{"name": "id", "type": "Row Number"}],
        "num_rows": MAX_ROWS + 1,
        "format": "CSV",
    })
    assert response.status_code == 400


def test_generate_returns_csv_attachment(client, schema):
    schema["format"] = "CSV"
    response = client.post("/generate", json=schema)
    assert response.status_code == 200
    assert response.mimetype == "text/csv"
    assert "attachment" in response.headers["Content-Disposition"]


def test_generate_returns_excel_with_xlsx_extension(client, schema):
    schema["format"] = "EXCEL"
    response = client.post("/generate", json=schema)
    assert response.status_code == 200
    assert ".xlsx" in response.headers["Content-Disposition"]


# ---- Template engine ------------------------------------------------------

def test_render_template_substitutes_simple_field():
    assert render_template_value("hello {{name}}", {"name": "world"}) == "hello world"


def test_render_template_with_filters():
    row = {"first": "Jane", "last": "Doe"}
    out = render_template_value("{{first|lower}}.{{last|lower}}@example.com", row)
    assert out == "jane.doe@example.com"


def test_render_template_chained_filters():
    assert render_template_value("{{name|upper|trim}}", {"name": "  alice "}) == "ALICE"


def test_render_template_initial_and_slug():
    assert render_template_value("{{name|initial}}", {"name": "alice"}) == "A"
    assert render_template_value("{{name|slug}}", {"name": "Hello World!"}) == "hello-world"


def test_render_template_unknown_field_renders_empty():
    assert render_template_value("X={{missing}}Y", {}) == "X=Y"


def test_render_template_handles_none():
    assert render_template_value("{{x|upper}}", {"x": None}) == ""


def test_template_field_resolves_after_others():
    schema = {
        "fields": [
            {"name": "first", "type": "First Name"},
            {"name": "last", "type": "Last Name"},
            {"name": "email", "type": "Template",
             "template": "{{first|lower}}.{{last|lower}}@example.com"},
        ],
        "num_rows": 5,
        "format": "JSON",
    }
    data = generate_data(schema)
    for row in data:
        assert row["email"].endswith("@example.com")
        first = row["first"].lower()
        last = row["last"].lower()
        assert row["email"] == f"{first}.{last}@example.com"


def test_template_field_with_blank_percentage():
    schema = {
        "fields": [
            {"name": "first", "type": "First Name"},
            {"name": "alias", "type": "Template",
             "template": "{{first|upper}}", "blank_percentage": 100},
        ],
        "num_rows": 8,
        "format": "JSON",
    }
    data = generate_data(schema)
    assert all(row["alias"] is None for row in data)


# ---- Schema inference: SQL ------------------------------------------------

def test_infer_from_sql_extracts_columns():
    sql = """
    CREATE TABLE users (
        id INTEGER PRIMARY KEY,
        first_name VARCHAR(100),
        email VARCHAR(255),
        created_at TIMESTAMP,
        balance DECIMAL(10,2),
        active BOOLEAN
    );
    """
    fields = infer_from_sql(sql)
    by_name = {f["name"]: f for f in fields}
    assert by_name["id"]["type"] == "Row Number"
    assert by_name["first_name"]["type"] == "First Name"
    assert by_name["email"]["type"] == "Email Address"
    assert by_name["created_at"]["type"] == "Date"
    assert by_name["balance"]["type"] == "Decimal"
    assert by_name["active"]["type"] == "Custom List"
    assert by_name["active"]["values"] == ["true", "false"]


def test_infer_from_sql_handles_multi_param_numeric_type():
    # Regression: _COLUMN_LINE_RE used to allow at most two numeric params,
    # so a type like NUMERIC(18,4,2) would fail to match the column line.
    sql = """
    CREATE TABLE t (
        id INT,
        amount NUMERIC(18,4,2)
    );
    """
    fields = infer_from_sql(sql)
    assert [f["name"] for f in fields] == ["id", "amount"]


def test_infer_from_sql_skips_constraints():
    sql = """
    CREATE TABLE t (
        id INT,
        name TEXT,
        PRIMARY KEY (id),
        UNIQUE (name)
    );
    """
    fields = infer_from_sql(sql)
    assert [f["name"] for f in fields] == ["id", "name"]


def test_infer_from_sql_rejects_garbage():
    with pytest.raises(SchemaError):
        infer_from_sql("this is not sql")


# ---- Schema inference: JSON -----------------------------------------------

def test_infer_from_json_object():
    src = '{"id": 1, "name": "Alice", "email": "a@b.com", "born": "1990-01-01"}'
    fields = infer_from_json(src)
    by_name = {f["name"]: f for f in fields}
    assert by_name["id"]["type"] == "Row Number"
    assert by_name["name"]["type"] == "Full Name"
    assert by_name["email"]["type"] == "Email Address"
    assert by_name["born"]["type"] == "Date"


def test_infer_from_json_array_uses_first_item():
    src = '[{"price": 9.99, "qty": 3}, {"price": 1.0, "qty": 1}]'
    fields = infer_from_json(src)
    by_name = {f["name"]: f for f in fields}
    assert by_name["price"]["type"] == "Decimal"
    assert by_name["qty"]["type"] == "Number"


def test_infer_from_json_bool_to_custom_list():
    fields = infer_from_json('{"active": true}')
    assert fields[0]["type"] == "Custom List"
    assert fields[0]["values"] == ["true", "false"]


def test_infer_from_json_detects_ipv4():
    fields = infer_from_json('{"client_addr": "192.168.1.1"}')
    assert fields[0]["type"] == "IP Address v4"


def test_infer_from_json_rejects_invalid():
    with pytest.raises(SchemaError):
        infer_from_json("not json")


def test_infer_from_json_rejects_empty_array():
    with pytest.raises(SchemaError):
        infer_from_json("[]")


# ---- Schema inference: TypeScript -----------------------------------------

def test_infer_from_typescript_interface():
    src = """
    interface User {
        id: number;
        email: string;
        firstName: string;
        active: boolean;
        createdAt: Date;
    }
    """
    fields = infer_from_typescript(src)
    by_name = {f["name"]: f for f in fields}
    assert by_name["id"]["type"] == "Row Number"
    assert by_name["email"]["type"] == "Email Address"
    assert by_name["firstName"]["type"] == "First Name"
    assert by_name["active"]["type"] == "Custom List"
    assert by_name["createdAt"]["type"] == "Date"


def test_infer_from_typescript_handles_type_alias():
    src = "type Product = { id: number; name: string; price: number; };"
    fields = infer_from_typescript(src)
    assert [f["name"] for f in fields] == ["id", "name", "price"]


def test_infer_from_typescript_optional_fields():
    src = "interface X { foo?: string; bar: number; }"
    fields = infer_from_typescript(src)
    assert [f["name"] for f in fields] == ["foo", "bar"]


def test_infer_from_typescript_rejects_garbage():
    with pytest.raises(SchemaError):
        infer_from_typescript("function foo() {}")


# ---- /infer-schema endpoint -----------------------------------------------

def test_infer_endpoint_sql(client):
    response = client.post("/infer-schema", json={
        "kind": "sql",
        "source": "CREATE TABLE t (id INT, email VARCHAR(100));",
    })
    assert response.status_code == 200
    fields = response.get_json()["fields"]
    assert any(f["name"] == "email" and f["type"] == "Email Address" for f in fields)


def test_infer_endpoint_json(client):
    response = client.post("/infer-schema", json={
        "kind": "json",
        "source": '{"name": "x", "email": "a@b.com"}',
    })
    assert response.status_code == 200
    fields = response.get_json()["fields"]
    assert {f["name"] for f in fields} == {"name", "email"}


def test_infer_endpoint_typescript(client):
    response = client.post("/infer-schema", json={
        "kind": "typescript",
        "source": "interface X { name: string; }",
    })
    assert response.status_code == 200


def test_infer_endpoint_rejects_unknown_kind(client):
    response = client.post("/infer-schema", json={"kind": "yaml", "source": "foo"})
    assert response.status_code == 400


def test_infer_endpoint_rejects_empty_source(client):
    response = client.post("/infer-schema", json={"kind": "sql", "source": "  "})
    assert response.status_code == 400


def test_infer_endpoint_rejects_oversize_source(client):
    response = client.post("/infer-schema", json={
        "kind": "json",
        "source": "x" * 40000,
    })
    assert response.status_code == 400


# ---- Hardening: XML name sanitizer + field-level validation ---------------

def test_sanitize_xml_name_replaces_invalid_chars():
    assert _sanitize_xml_name("user name") == "user_name"
    assert _sanitize_xml_name("col<>&") == "col___"
    assert _sanitize_xml_name("") == "_field"


def test_sanitize_xml_name_prepends_underscore_for_digit_start():
    assert _sanitize_xml_name("1col") == "_1col"


def test_format_xml_accepts_field_names_starting_with_digit():
    # Regression: this used to raise ValueError → 500.
    data = [{"1col": "value"}]
    out = format_xml(data)
    assert "<_1col>value</_1col>" in out


def test_format_xml_accepts_field_names_with_spaces():
    data = [{"col with space": "v"}]
    out = format_xml(data)
    assert "<col_with_space>v</col_with_space>" in out


def test_validate_schema_rejects_field_without_name():
    with pytest.raises(SchemaError):
        _validate_schema({"fields": [{"type": "First Name"}], "num_rows": 1, "format": "CSV"})


def test_validate_schema_rejects_field_without_type():
    with pytest.raises(SchemaError):
        _validate_schema({"fields": [{"name": "x"}], "num_rows": 1, "format": "CSV"})


def test_validate_schema_rejects_non_dict_field():
    with pytest.raises(SchemaError):
        _validate_schema({"fields": ["not a dict"], "num_rows": 1, "format": "CSV"})


def test_validate_schema_rejects_custom_list_without_values():
    with pytest.raises(SchemaError):
        _validate_schema({
            "fields": [{"name": "x", "type": "Custom List"}],
            "num_rows": 1, "format": "CSV",
        })


def test_validate_schema_rejects_template_without_template():
    with pytest.raises(SchemaError):
        _validate_schema({
            "fields": [{"name": "x", "type": "Template"}],
            "num_rows": 1, "format": "CSV",
        })


def test_validate_schema_rejects_too_many_fields():
    fields = [{"name": f"c{i}", "type": "First Name"} for i in range(MAX_FIELDS + 1)]
    with pytest.raises(SchemaError):
        _validate_schema({"fields": fields, "num_rows": 1, "format": "CSV"})


def test_validate_schema_rejects_bad_blank_percentage():
    with pytest.raises(SchemaError):
        _validate_schema({
            "fields": [{"name": "x", "type": "First Name", "blank_percentage": 150}],
            "num_rows": 1, "format": "CSV",
        })


def test_endpoint_rejects_malformed_field(client):
    response = client.post("/preview", json={
        "fields": [{"type": "First Name"}],  # no name
        "num_rows": 1, "format": "CSV",
    })
    assert response.status_code == 400
    assert "name" in response.get_json()["error"]


def test_endpoint_returns_413_for_oversize_request(client):
    # Build a payload larger than MAX_CONTENT_LENGTH.
    huge_source = "x" * (MAX_CONTENT_LENGTH + 100)
    response = client.post(
        "/infer-schema",
        data=f'{{"kind":"json","source":"{huge_source}"}}',
        content_type="application/json",
    )
    assert response.status_code == 413


# ---- Review follow-ups ------------------------------------------------------

def _schema(fields, num_rows=1, fmt="CSV", **extra):
    return {"fields": fields, "num_rows": num_rows, "format": fmt, **extra}


def test_format_csv_can_skip_formula_escaping():
    data = [{"phone": "+1-202-555-0100", "f": "=1+1"}]
    lines = format_csv(data, escape_formulas=False).strip().splitlines()
    assert lines[1] == "+1-202-555-0100,=1+1"


def test_generate_endpoint_honours_csv_formula_escape_flag(client):
    fields = [{"name": "v", "type": "Custom List", "values": ["+1-202-555-0100"]}]
    raw = client.post("/generate", json=_schema(fields, csv_formula_escape=False))
    escaped = client.post("/generate", json=_schema(fields))
    assert raw.get_data(as_text=True).splitlines()[1] == "+1-202-555-0100"
    assert escaped.get_data(as_text=True).splitlines()[1] == "'+1-202-555-0100"


def test_validate_schema_rejects_non_bool_csv_formula_escape():
    with pytest.raises(SchemaError):
        _validate_schema(_schema([{"name": "a", "type": "City"}], csv_formula_escape="no"))


@pytest.mark.parametrize(
    "template",
    [
        "{{a|" + " " * 996,       # unclosed filter chain followed by whitespace
        "{{" * 500,               # flood of opening braces
        "{{" + "a" * 998,         # unclosed field name
        "{{a| x" * 166,
    ],
)
def test_template_parsing_is_fast_on_pathological_input(template):
    # Regression: the old token pattern backtracked catastrophically —
    # a ~2 KB unclosed token took ~26 s to render, once per generated row.
    import time

    start = time.perf_counter()
    compile_template(template)
    assert time.perf_counter() - start < 0.5


def test_template_keeps_literal_braces_around_tokens():
    assert render_template_value("{{{a}}}", {"a": "X"}) == "{X}"


def test_template_with_empty_filter_chain_renders_value():
    assert render_template_value("{{a|}}", {"a": "X"}) == "X"


def test_template_ignores_unknown_filters_and_whitespace():
    assert render_template_value("{{ a | upper | nope }}", {"a": "x"}) == "X"


def test_validate_schema_rejects_overlong_template():
    field = {"name": "t", "type": "Template", "template": "x" * (MAX_TEMPLATE_LENGTH + 1)}
    with pytest.raises(SchemaError):
        _validate_schema(_schema([field]))


def test_validate_schema_normalizes_numeric_string_blank_percentage():
    # Regression: "50" passed validation but the generator then divided a
    # string by 100 and returned a 500.
    schema = _validate_schema(_schema([{"name": "a", "type": "City", "blank_percentage": "50"}]))
    assert schema["fields"][0]["blank_percentage"] == 50
    generate_data(schema)  # must not raise


def test_validate_schema_rejects_unknown_type():
    with pytest.raises(SchemaError, match="unknown type"):
        _validate_schema(_schema([{"name": "a", "type": "Not A Type"}]))


def test_validate_schema_rejects_duplicate_names():
    with pytest.raises(SchemaError, match="Duplicate"):
        _validate_schema(_schema([{"name": "a", "type": "City"}, {"name": "a", "type": "Country"}]))


def test_validate_schema_rejects_non_scalar_custom_list_values():
    field = {"name": "a", "type": "Custom List", "values": [{"nested": 1}]}
    with pytest.raises(SchemaError):
        _validate_schema(_schema([field]))


def test_validate_schema_rejects_too_many_cells():
    fields = [{"name": f"c{i}", "type": "City"} for i in range(20)]
    rows = MAX_CELLS // 20 + 1
    if rows > MAX_ROWS:
        pytest.skip("MAX_ROWS is below the cell cap for this field count")
    with pytest.raises(SchemaError, match="cells"):
        _validate_schema(_schema(fields, num_rows=rows))


def test_format_sql_states_target_dialect():
    assert "-- Dialect: SQLite / PostgreSQL" in format_sql([{"a": 1}])


@pytest.mark.parametrize(
    ("name", "expected"),
    [
        ("paid", "(fallback)"),
        ("is_valid", "(fallback)"),
        ("uuid", "(fallback)"),
        ("message", "(fallback)"),
        ("language", "(fallback)"),
        ("hotel_name", "(fallback)"),
        ("capacity", "(fallback)"),
        ("destination", "(fallback)"),
        ("account_name", "(fallback)"),
        ("cancelled_at", "Date"),
        ("publishedAt", "Date"),
        ("discount", "Decimal"),
        ("view_count", "Number"),
        ("userId", "Row Number"),
        ("IPAddress", "IP Address v4"),
        ("e_mail", "Email Address"),
        ("phone_number", "Phone Number"),
        ("lastName", "Last Name"),
        ("display_name", "Full Name"),
        ("years_experience", "Number"),
        ("unit_prices", "Decimal"),
        ("emails", "Email Address"),
        ("cities", "City"),
        ("address", "(fallback)"),
        ("status", "(fallback)"),
    ],
)
def test_name_heuristics_match_whole_words(name, expected):
    assert _name_to_type(name, "(fallback)") == expected
